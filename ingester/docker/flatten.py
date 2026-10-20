"""Flatten a registered document into the fragment table (F16) that every document citation points at.

Usage: flatten.py <engagement> <Dnnn>

Reads engagements/<engagement>/documents/Dnnn.md for the file name and SHA-256, finds the file under
documents/unprocessed/ or documents/processed/, checks its SHA-256, and writes three files to
sessions/Dnnn/:

  Dnnn.fragments.tsv  unit fragment location section kind text, one row per fragment (F16)
  Dnnn.sections.tsv   section slug fragments first-unit last-unit state, one row per section
  Dnnn.meta.tsv       key value: format, sha256, author, date, title from the file's own properties

xlsx: a unit is a row that has any text, numbered across the workbook; a fragment is a cell, numbered by
its column (A is 1); the location is Sheet!C17; the section is the sheet; state is visible, hidden or
veryHidden. pptx: a unit is a slide; a fragment is a paragraph or table cell, then each paragraph of the
speaker notes; the location is "slide 12" or "slide 12 notes"; the sections are Slides and Notes. pdf: a
unit is a page; a fragment is a line of text, then each table cell; the location is "p. 7"; the one
section is Pages.

Text has its whitespace collapsed to single spaces, so a quote copied from the table matches the check.
Runs in the ingester's Docker image (make flatten); writes nothing outside sessions/Dnnn/.
"""

import datetime
import hashlib
import os
import re
import sys

ENGAGEMENTS = os.environ.get("ENGAGEMENTS_DIR", "/engagements")
SESSIONS = os.environ.get("SESSIONS_DIR", "/sessions")


def die(msg):
    sys.stderr.write("error: %s\n" % msg)
    sys.exit(1)


def frontmatter(path):
    out = {}
    with open(path, encoding="utf-8") as f:
        lines = f.read().split("\n")
    if not lines or lines[0] != "---":
        die("%s has no frontmatter" % path)
    for line in lines[1:]:
        if line == "---":
            break
        if ": " in line:
            k, v = line.split(": ", 1)
            out[k] = v.strip()
    return out


def clean(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime.datetime):
        value = value.date() if value.time() == datetime.time(0, 0) else value
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def plain_date(value):
    if not isinstance(value, (datetime.date, datetime.datetime)):
        return ""
    return "%d %s %d" % (value.day, value.strftime("%B"), value.year)


def slugify(name, used):
    base = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-") or "section"
    slug, n = base, 2
    while slug in used:
        slug, n = "%s-%d" % (base, n), n + 1
    used.add(slug)
    return slug


def flatten_xlsx(path):
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    rows, sections, unit = [], [], 0
    for ws in wb.worksheets:
        first = last = None
        count = 0
        for row in ws.iter_rows():
            cells = [(c.column, c.coordinate, clean(c.value)) for c in row if clean(c.value) != ""]
            if not cells:
                continue
            unit += 1
            first = unit if first is None else first
            last = unit
            for col, coord, text in cells:
                rows.append((unit, col, "%s!%s" % (ws.title.replace("|", "/"), coord), ws.title, "cell", text))
                count += 1
        sections.append((ws.title, count, first or "", last or "", ws.sheet_state))
    p = wb.properties
    meta = {"author": clean(p.creator), "date": plain_date(p.modified or p.created), "title": clean(p.title)}
    return rows, sections, meta


def flatten_pptx(path):
    from pptx import Presentation

    prs = Presentation(path)
    rows, counts, spans = [], {"Slides": 0, "Notes": 0}, {"Slides": [None, None], "Notes": [None, None]}

    def add(unit, frag, loc, section, kind, text):
        rows.append((unit, frag, loc, section, kind, text))
        counts[section] += 1
        span = spans[section]
        span[0] = unit if span[0] is None else span[0]
        span[1] = unit

    def shapes(group):
        for sh in group:
            if sh.shape_type == 6:  # group shape
                yield from shapes(sh.shapes)
            else:
                yield sh

    for i, slide in enumerate(prs.slides, 1):
        frag = 0
        for sh in shapes(slide.shapes):
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    text = clean("".join(r.text for r in para.runs))
                    if text:
                        frag += 1
                        add(i, frag, "slide %d" % i, "Slides", "para", text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    for c in r.cells:
                        text = clean(c.text)
                        if text:
                            frag += 1
                            add(i, frag, "slide %d" % i, "Slides", "table", text)
        if slide.has_notes_slide:
            for para in slide.notes_slide.notes_text_frame.paragraphs:
                text = clean("".join(r.text for r in para.runs))
                if text:
                    frag += 1
                    add(i, frag, "slide %d notes" % i, "Notes", "notes", text)
    sections = [(s, counts[s], spans[s][0] or "", spans[s][1] or "", "visible") for s in ("Slides", "Notes") if counts[s]]
    p = prs.core_properties
    meta = {"author": clean(p.author), "date": plain_date(p.modified or p.created), "title": clean(p.title)}
    return rows, sections, meta


def flatten_pdf(path):
    import pdfplumber

    rows, count, first, last = [], 0, None, None
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            frag = 0
            for line in (page.extract_text() or "").split("\n"):
                text = clean(line)
                if text:
                    frag += 1
                    rows.append((i, frag, "p. %d" % i, "Pages", "line", text))
            for table in page.extract_tables() or []:
                for r in table:
                    for c in r:
                        text = clean(c)
                        if text:
                            frag += 1
                            rows.append((i, frag, "p. %d" % i, "Pages", "table", text))
            if frag:
                count += frag
                first = i if first is None else first
                last = i
        info = pdf.metadata or {}
    date = ""
    m = re.match(r"D:(\d{4})(\d{2})(\d{2})", str(info.get("ModDate") or info.get("CreationDate") or ""))
    if m:
        date = plain_date(datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3))))
    meta = {"author": clean(info.get("Author")), "date": date, "title": clean(info.get("Title"))}
    return rows, [("Pages", count, first or "", last or "", "visible")], meta


def main():
    if len(sys.argv) != 3:
        die("usage: flatten.py <engagement> <Dnnn>")
    eng, did = sys.argv[1], sys.argv[2]
    if not re.match(r"^D\d{3}$", did):
        die("%s is not a document id" % did)
    base = os.path.join(ENGAGEMENTS, eng)
    reg = os.path.join(base, "documents", did + ".md")
    if not os.path.isfile(reg):
        die("%s is not registered: no %s" % (did, reg))
    fm = frontmatter(reg)
    src = None
    for where in ("unprocessed", "processed"):
        p = os.path.join(base, "documents", where, fm.get("file", ""))
        if os.path.isfile(p):
            src = p
            break
    if src is None:
        die("document file %s not found under documents/" % fm.get("file"))
    with open(src, "rb") as f:
        sha = hashlib.sha256(f.read()).hexdigest()
    if sha != fm.get("sha256"):
        die("SHA-256 of %s does not match documents/%s.md" % (src, did))
    ext = os.path.splitext(src)[1].lower()
    handler = {".xlsx": flatten_xlsx, ".pptx": flatten_pptx, ".pdf": flatten_pdf}.get(ext)
    if handler is None:
        die("%s is not .xlsx, .pptx or .pdf" % src)
    rows, sections, meta = handler(src)

    out = os.path.join(SESSIONS, did)
    os.makedirs(out, exist_ok=True)

    def tsv(name, header, data):
        tmp = os.path.join(out, name + ".tmp")
        with open(tmp, "w", encoding="utf-8", newline="\n") as f:
            f.write("\t".join(header) + "\n")
            for r in data:
                f.write("\t".join(clean(v) if not isinstance(v, int) else str(v) for v in r) + "\n")
        os.replace(tmp, os.path.join(out, name))

    used = set()
    tsv(did + ".fragments.tsv", ["unit", "fragment", "location", "section", "kind", "text"], sorted(rows, key=lambda r: (r[0], r[1])))
    tsv(did + ".sections.tsv", ["section", "slug", "fragments", "first-unit", "last-unit", "state"],
        [(s, slugify(s, used), n, a, b, st) for s, n, a, b, st in sections])
    meta.update({"format": ext[1:], "sha256": sha, "file": fm.get("file", "")})
    tsv(did + ".meta.tsv", ["key", "value"], sorted(meta.items()))
    print("flattened %s: %d fragment(s) in %d section(s) from %s" % (did, len(rows), len(sections), os.path.basename(src)))


if __name__ == "__main__":
    main()
